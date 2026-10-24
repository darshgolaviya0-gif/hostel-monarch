"""
HOSTEL MONARCH — Seminar III PPT (Redesigned Version 2)
Tapi Diploma Engineering College | Computer Department
Minor Project — DI05000341

Install: pip install python-pptx
Run:     python hostel_monarch_seminar3_v2.py
Output:  HOSTEL_MONARCH_Seminar_3_Redesigned.pptx

The deck is intentionally honest: because development has not started, planned
items are labelled PROPOSED or PENDING. Replace placeholders/results only after
real implementation and testing.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

# ───────────────────────────── EDITABLE DATA ──────────────────────────────
PROJECT = {
    "title": "HOSTEL MONARCH",
    "tagline": "Smart and Interactive Hostel Management System",
    "college": "TAPI DIPLOMA ENGINEERING COLLEGE",
    "department": "COMPUTER DEPARTMENT",
    "course": "Minor Project — DI05000341",
    "seminar": "SEMINAR–III",
    "year": "Academic Year 2024–25",  # update if needed
    "guide": "Computer Department",   # replace with guide's name
}

# Rudra Bhanderi has been removed as requested.
TEAM = [
    ("Aryan Dodia", "246470307003"),
    ("Darsh Golaviya", "246470307018"),
    ("Divyam Bhimani", "246470307023"),
    ("Krish Mendpara", "246470307059"),
]

# Proposed stack — finalise before development and then remove “Proposed”.
STACK = [
    ("Frontend", "HTML5 · CSS3 · JavaScript · Bootstrap 5"),
    ("Backend", "PHP 8 — Proposed"),
    ("Database", "MySQL 8"),
    ("Server", "Apache through XAMPP"),
    ("IDE", "Visual Studio Code"),
    ("Tools", "Git · GitHub · phpMyAdmin · Postman"),
    ("QR", "Open-source PHP QR library"),
    ("Hosting", "Local server; LAMP hosting after testing"),
]

ASSETS = {
    "login": "assets/login.png",
    "student": "assets/student_dashboard.png",
    "admin": "assets/admin_dashboard.png",
    "allocation": "assets/room_allocation.png",
    "qr": "assets/leave_qr.png",
    "demo": "assets/demo.png",
}

OUTPUT = "HOSTEL_MONARCH_Seminar_3_Redesigned.pptx"

# ─────────────────────────────── THEME ────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

INK = RGBColor(22, 29, 43)
BODY = RGBColor(75, 84, 99)
GREEN = RGBColor(7, 72, 57)
ACCENT = RGBColor(17, 170, 124)
MINT = RGBColor(228, 247, 240)
PALE = RGBColor(246, 248, 248)
LINE = RGBColor(221, 228, 227)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(54, 112, 181)
AMBER = RGBColor(230, 158, 48)
RED = RGBColor(195, 63, 63)
FONT = "Aptos"


def add_rect(slide, x, y, w, h, fill, radius=False, outline=None, width=1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = outline or fill; sh.line.width = Pt(width)
    return sh


def add_text(slide, text, x, y, w, h, size=16, color=BODY, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=.03):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold
    p.font.color.rgb = color
    return sh


def add_bullets(slide, items, x, y, w, h, size=14, color=BODY, spacing=7):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item; p.font.name = FONT; p.font.size = Pt(size)
        p.font.color.rgb = color; p.space_after = Pt(spacing)
    return sh


def base_slide(title, subtitle=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    add_text(s, title, .72, .42, 11.7, .43, 24, INK, True)
    if subtitle:
        add_text(s, subtitle, .72, .91, 11.55, .28, 10.5, ACCENT, True)
    add_rect(s, .72, 7.03, .55, .035, ACCENT)
    add_text(s, "HOSTEL MONARCH  ·  SEMINAR III", 1.42, 6.98, 3.5, .15, 7.5, BODY)
    add_text(s, f"{len(prs.slides):02d}", 12.25, 6.98, .35, .15, 7.5, BODY, False, PP_ALIGN.RIGHT)
    return s


def card(slide, x, y, w, h, heading, body, accent=ACCENT, number=None):
    add_rect(slide, x, y, w, h, PALE, True, LINE)
    add_rect(slide, x, y, .06, h, accent, True, accent)
    if number is not None:
        add_rect(slide, x+.22, y+.18, .46, .46, accent, True)
        add_text(slide, str(number), x+.22, y+.31, .46, .16, 10, WHITE, True, PP_ALIGN.CENTER)
        hx = x+.82
    else:
        hx = x+.25
    add_text(slide, heading, hx, y+.18, w-(hx-x)-.20, .25, 12.5, GREEN, True)
    add_text(slide, body, x+.25, y+.60, w-.50, h-.72, 10.6, BODY)


def pill(slide, text, x, y, w, color=AMBER):
    add_rect(slide, x, y, w, .34, color, True)
    add_text(slide, text, x+.06, y+.085, w-.12, .15, 8.5, WHITE, True, PP_ALIGN.CENTER)


def add_table(slide, data, x, y, w, h, widths=None, size=9.2):
    rows, cols = len(data), len(data[0])
    tbl = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for i, val in enumerate(widths): tbl.columns[i].width = Inches(val)
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c); cell.text = str(data[r][c])
            cell.margin_left = cell.margin_right = Inches(.10)
            cell.margin_top = cell.margin_bottom = Inches(.05)
            cell.fill.solid(); cell.fill.fore_color.rgb = GREEN if r == 0 else (WHITE if r % 2 else PALE)
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT; p.font.size = Pt(size)
                p.font.bold = r == 0; p.font.color.rgb = WHITE if r == 0 else BODY
    return tbl


def arrow(slide, x1, y1, x2, y2, color=ACCENT, width=1.7):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width); ln.line.end_arrowhead = True
    return ln


def screenshot(slide, path, x, y, w, h, label):
    p = Path(path)
    if p.exists():
        slide.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        add_rect(slide, x, y, w, h, PALE, True, LINE)
        add_rect(slide, x+.18, y+.18, w-.36, h-.36, WHITE, True, LINE)
        add_text(slide, "SCREENSHOT", x+.4, y+h/2-.42, w-.8, .25, 12, ACCENT, True, PP_ALIGN.CENTER)
        add_text(slide, label, x+.4, y+h/2-.03, w-.8, .30, 11, INK, True, PP_ALIGN.CENTER)
        add_text(slide, f"Add file: {path}", x+.4, y+h/2+.38, w-.8, .20, 8.5, BODY, False, PP_ALIGN.CENTER)

# ─────────────────────────────── SLIDES ───────────────────────────────────
# 1 Title
s = prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb = GREEN
add_rect(s, 0, 0, .22, 7.5, ACCENT)
add_text(s, f"{PROJECT['college']}  ·  {PROJECT['department']}", .95, .78, 9, .24, 9.5, ACCENT, True)
add_text(s, PROJECT["title"], .95, 1.62, 10.5, .66, 34, WHITE, True)
add_text(s, PROJECT["tagline"], .95, 2.42, 10.5, .34, 16, WHITE)
add_rect(s, .95, 2.98, 2.2, .04, ACCENT)
add_text(s, f"{PROJECT['seminar']}  ·  {PROJECT['course']}", .95, 3.42, 6.0, .28, 11, ACCENT, True)
add_text(s, PROJECT["year"], .95, 3.82, 4.0, .26, 12.5, WHITE)
add_text(s, "PROJECT TEAM", 7.05, 3.42, 2.0, .24, 10, ACCENT, True)
add_text(s, "\n".join(f"{n} — {e}" for n,e in TEAM), 7.05, 3.80, 4.8, 1.30, 12.5, WHITE)
pill(s, "CURRENT STAGE: DESIGN · IMPLEMENTATION PENDING", .95, 6.40, 3.85, AMBER)

# 2 Index
s = base_slide("Presentation Index", "Structured according to the Seminar–III guidelines")
items = ["Project Title", "Project Design & UML", "Software Technologies", "Final Hardware Assembly",
         "Software Implementation", "Testing Methodology", "System Testing", "Demonstration",
         "Verification & Performance", "Conclusion"]
for i, item in enumerate(items):
    col, row = i % 2, i // 2; x = .75 + col*6.12; y = 1.42 + row*.98
    add_rect(s, x, y, 5.72, .68, PALE, True, LINE)
    add_text(s, f"{i+1:02d}", x+.20, y+.20, .42, .18, 10, ACCENT, True)
    add_text(s, item, x+.82, y+.17, 4.60, .25, 12.2, INK, True)

# 3 Project overview
s = base_slide("Project Title & Overview", "One platform for the complete hostel service lifecycle")
card(s,.75,1.42,3.72,1.62,"Problem","Paper records, repeated office visits, disconnected workflows and limited visibility create delay and errors.",RED,"01")
card(s,4.80,1.42,3.72,1.62,"Solution","A secure role-based web platform connecting students, management, security, maintenance and mess staff.",ACCENT,"02")
card(s,8.85,1.42,3.72,1.62,"Value","Faster service, transparent records, improved traceability and reliable data for hostel decisions.",BLUE,"03")
add_text(s,"CORE MODULES",.75,3.55,2.0,.25,12,GREEN,True)
modules=["Authentication","Profile & Admission","Room Allocation","Fees & Receipts","Complaints",
         "Leave & QR","Visitor & Parcel","Mess Services","Notices","Reports & Audit"]
for i,m in enumerate(modules):
    col=i%5; row=i//5; x=.75+col*2.42; y=4.02+row*.93
    add_rect(s,x,y,2.12,.60,MINT,True,ACCENT)
    add_text(s,m,x+.10,y+.19,1.92,.20,9.7,GREEN,True,PP_ALIGN.CENTER)

# 4 Honest status
s = base_slide("Project Status", "Clear baseline before development begins")
pill(s,"DESIGN / PRE-DEVELOPMENT",.75,1.42,2.65,AMBER)
rows=[["Work Area","Current Status","Next Action"],
      ["Requirements & scope","Completed","Freeze minimum viable product"],
      ["UML & database design","Draft prepared","Technical review and normalisation"],
      ["Technology stack","Proposed","Team must confirm one backend stack"],
      ["Frontend","Not started","Build shared responsive components"],
      ["Backend & database","Not started","Implement authentication and schema"],
      ["Testing & demo","Pending","Execute only after working build"]]
add_table(s,rows,.75,2.00,11.85,3.95,[3.05,2.55,6.25],10)
add_text(s,"No implementation, test or performance result is claimed in this presentation.",.78,6.40,11.5,.28,10.5,RED,True)

# 5 Use case
s = base_slide("Project Design — Use-Case Diagram", "Primary actions available to each role")
# central system
add_rect(s,3.25,1.38,6.85,5.18,WHITE,True,LINE)
add_text(s,"HOSTEL MONARCH",5.10,1.60,3.15,.24,11,GREEN,True,PP_ALIGN.CENTER)
actors=[("STUDENT",.65,2.00),("WARDEN / ADMIN",.50,4.78),("SECURITY",11.05,2.00),("SERVICE STAFF",10.78,4.78)]
for name,x,y in actors:
    add_rect(s,x,y,1.72,.58,PALE,True,LINE); add_text(s,name,x+.06,y+.19,1.60,.18,8.8,INK,True,PP_ALIGN.CENTER)
uses=[("Register / Login",3.65,2.08),("Apply / Upload",5.48,2.08),("Allocate Bed",7.31,2.08),
      ("Fees / Receipt",3.65,3.18),("Complaint",5.48,3.18),("Leave / QR",7.31,3.18),
      ("Gate Scan",3.65,4.28),("Mess / Notice",5.48,4.28),("Reports / Audit",7.31,4.28),
      ("Visitor / Parcel",5.48,5.38)]
for text,x,y in uses:
    add_rect(s,x,y,1.43,.55,MINT,True,ACCENT); add_text(s,text,x+.05,y+.18,1.33,.17,8.5,GREEN,True,PP_ALIGN.CENTER)
# associations kept light and uncluttered
for coords in [(2.37,2.29,3.65,2.35),(2.37,2.29,5.48,2.35),(2.37,2.29,3.65,3.45),
               (2.22,5.06,7.31,2.35),(2.22,5.06,7.31,4.55),(11.05,2.29,5.08,4.55),
               (10.78,5.06,4.98,3.45),(10.78,5.06,6.91,4.55)]: arrow(s,*coords,LINE,1)

# 6 Activity flow
s = base_slide("Project Design — Activity Flow", "Admission, allocation and payment journey")
flow=[("START",GREEN),("REGISTER",ACCENT),("PROFILE",ACCENT),("APPLICATION",ACCENT),
      ("VERIFY",BLUE),("ALLOCATE",BLUE),("INVOICE",AMBER),("PAYMENT",AMBER),("RECEIPT",ACCENT),("END",GREEN)]
for i,(label,col) in enumerate(flow):
    x=.48+i*1.27; w=.88 if i in (0,9) else 1.05
    add_rect(s,x,2.55,w,.62,col if i in (0,9) else PALE,True,col)
    add_text(s,label,x+.04,2.78,w-.08,.17,8.6,WHITE if i in (0,9) else INK,True,PP_ALIGN.CENTER)
    if i<9: arrow(s,x+w,2.86,x+1.23,2.86,ACCENT,1.5)
add_text(s,"CONTROLLED EXCEPTION PATHS",.75,4.02,3.0,.25,11.5,GREEN,True)
card(s,.75,4.48,3.72,1.38,"Document rejected","Return the application with remarks for correction and resubmission.",RED)
card(s,4.80,4.48,3.72,1.38,"Bed unavailable","Reject conflicting allocation and require another available bed.",AMBER)
card(s,8.85,4.48,3.72,1.38,"Payment failed","Keep invoice unpaid; do not generate a successful receipt.",BLUE)

# 7 Sequence
s = base_slide("Project Design — Sequence Diagram", "Leave approval and QR gate verification")
names=["Student","Web UI","Server","Warden","Gate","Database"]
xs=[.55,2.70,4.85,7.00,9.15,11.30]
for name,x in zip(names,xs):
    add_rect(s,x,1.40,1.45,.48,GREEN,True); add_text(s,name,x+.04,1.56,1.37,.16,9,WHITE,True,PP_ALIGN.CENTER)
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x+.725),Inches(1.88),Inches(x+.725),Inches(6.42)); ln.line.color.rgb=LINE
messages=[(0,1,2.22,"Submit request"),(1,2,2.67,"Validate"),(2,5,3.12,"Save pending"),(2,3,3.57,"Approval request"),
          (3,2,4.02,"Approve / reject"),(2,5,4.47,"Update status"),(2,4,4.92,"Generate QR"),(4,2,5.37,"Verify scan"),
          (2,5,5.82,"Create gate log"),(2,1,6.27,"Return result")]
for a,b,y,msg in messages:
    arrow(s,xs[a]+.73,y,xs[b]+.73,y,ACCENT if b>a else BLUE,1.15)
    add_text(s,msg,min(xs[a],xs[b])+.78,y-.20,abs(xs[b]-xs[a])-.12,.16,7.7,BODY,False,PP_ALIGN.CENTER)

# 8 ER design
s = base_slide("Project Design — Database Model", "Core entities and technically appropriate cardinalities")
rels=[["Relationship","Cardinality","Reason"],
      ["User → Student","1 : 0..1","Not every user is a student"],
      ["Student → Application","1 : N","Retain application history"],
      ["Room → Bed","1 : N","A room contains multiple beds"],
      ["Bed → Allocation","1 : N","Retain allocation history"],
      ["Student → Allocation","1 : N","Support transfers/semesters"],
      ["Student → Invoice","1 : N","Multiple billing periods"],
      ["Invoice → Payment","1 : N","Support retry/partial history"],
      ["Leave Request → Gate Log","1 : N","Exit and re-entry events"]]
add_table(s,rels,.75,1.42,7.35,4.95,[3.15,1.30,2.90],9.3)
add_text(s,"KEY ENTITIES",8.48,1.48,1.8,.22,11,GREEN,True)
entities=["users / roles","students / applications","documents","rooms / beds","allocations","invoices / payments","complaints","leave / QR / gate logs","visitors / parcels","menus / notices / audit"]
for i,e in enumerate(entities):
    y=1.92+i*.43; add_rect(s,8.48,y,3.75,.31,MINT,True,MINT); add_text(s,e,8.65,y+.07,3.40,.13,8.8,GREEN,True)
add_text(s,"Rule: one bed may have only one active allocation at a time.",8.48,6.36,3.85,.25,9.3,RED,True)

# 9 technologies
s = base_slide("Software Technologies", "Proposed development environment — confirm before coding")
for i,(key,val) in enumerate(STACK):
    col=i%2; row=i//2; x=.75+col*6.12; y=1.42+row*1.13
    add_rect(s,x,y,5.72,.85,PALE,True,LINE)
    add_text(s,key.upper(),x+.20,y+.19,1.35,.18,8.8,ACCENT,True)
    add_text(s,val,x+1.63,y+.17,3.82,.31,10.8,INK,True)
pill(s,"STACK DECISION REQUIRED",4.88,6.20,3.55,AMBER)

# 10 hardware
s = base_slide("Final Hardware Assembly", "Not applicable as a custom assembly for this web-based project")
card(s,.75,1.42,5.75,1.75,"Academic prototype","No sensors, motors or microcontrollers are required. Hostel Monarch is a browser-based information system.",ACCENT,"✓")
card(s,6.85,1.42,5.75,1.75,"Required devices","Development computer, local/hosted server, user browser and phone/webcam for QR scanning.",BLUE,"D")
add_text(s,"LOGICAL DEPLOYMENT",.75,3.85,2.4,.22,11.5,GREEN,True)
boxes=[("Browser",.75),("Apache / PHP",3.72),("MySQL",6.69),("QR Camera",9.66)]
for i,(label,x) in enumerate(boxes):
    add_rect(s,x,4.45,2.25,.78,PALE,True,ACCENT); add_text(s,label,x+.10,4.73,2.05,.20,11,INK,True,PP_ALIGN.CENTER)
    if i<3: arrow(s,x+2.25,4.84,x+2.85,4.84,ACCENT,1.8)
add_text(s,"Future scope: a dedicated security-gate kiosk, RFID reader or biometric terminal.",.75,6.15,11.85,.30,10.5,BODY,False,PP_ALIGN.CENTER)

# 11 implementation architecture
s = base_slide("Software Implementation", "Planned architecture and development responsibilities")
layers=[("01","Presentation","Responsive pages, dashboards, forms and QR view",ACCENT),
        ("02","Application","Authentication, approval rules, allocation and reports",BLUE),
        ("03","Data Access","Prepared queries, transactions and reusable models",AMBER),
        ("04","Database","Relational tables, constraints, indexes and audit data",GREEN)]
for i,(n,h,b,col) in enumerate(layers):
    y=1.40+i*1.18; add_rect(s,1.05,y,11.20,.82,PALE,True,col); add_rect(s,1.05,y,.09,.82,col)
    add_text(s,n,1.40,y+.25,.38,.16,9,col,True)
    add_text(s,h,1.95,y+.20,2.05,.22,12,col,True)
    add_text(s,b,4.15,y+.20,7.55,.30,11,INK)
pill(s,"IMPLEMENTATION NOT STARTED",4.86,6.25,3.60,RED)

# 12 Frontend/backend
s = base_slide("Frontend & Backend Development", "Expected components and controls")
add_text(s,"FRONTEND",.75,1.42,2.0,.22,11.5,GREEN,True)
add_bullets(s,["Reusable navigation, cards, tables and forms","Responsive student and staff dashboards","Client-side usability validation","Loading, empty, success and error states","Accessible labels and status colours"],.78,1.86,5.45,3.55,12.6)
add_text(s,"BACKEND",6.75,1.42,2.0,.22,11.5,GREEN,True)
add_bullets(s,["Session-based authentication and role authorization","Server-side validation for every request","Allocation, approval and payment business rules","Prepared SQL statements and transactions","Audit logs, backups and controlled file uploads"],6.78,1.86,5.45,3.55,12.6)
add_rect(s,.75,5.65,11.85,.62,MINT,True,ACCENT)
add_text(s,"Request → Authenticate → Authorize → Validate → Apply Rule → Database Transaction → Response",1.05,5.86,11.25,.22,11,GREEN,True,PP_ALIGN.CENTER)
pill(s,"EVIDENCE PENDING",5.12,6.48,3.10,RED)

# 13 DB/API
s = base_slide("Database Development & API Integration", "Integrity controls and proposed internal endpoints")
api=[["Method","Endpoint","Purpose"],["POST","/auth/login","Authenticate and create session"],
     ["GET","/beds/available","Return available room beds"],["POST","/allocations","Create approved allocation"],
     ["POST","/complaints","Register complaint ticket"],["PATCH","/leaves/{id}","Approve or reject request"],
     ["POST","/gate/verify","Validate QR and create log"]]
add_table(s,api,.75,1.42,6.95,4.55,[.85,2.45,3.65],9.2)
add_text(s,"DATABASE CONTROLS",8.08,1.48,2.3,.22,11,GREEN,True)
add_bullets(s,["Primary and foreign keys","Unique active-bed rule","Password hashing","Prepared statements","Transactions and rollback","File type/size checks","Scheduled database backup"],8.08,1.90,4.25,3.75,11.5)
add_text(s,"Payment remains test-mode only; no bank/card credentials are stored.",.78,6.30,11.55,.26,9.8,RED,True)

# 14 methodology
s = base_slide("Testing Methodology", "Testing will begin as soon as executable modules are available")
tests=[("Functional","Business rules and expected outputs"),("UI","Layout, feedback and responsiveness"),
       ("Integration","Data exchange between modules"),("Database","CRUD, constraints and rollback"),
       ("Compatibility","Desktop and mobile browsers"),("Performance","Basic response and load checks")]
for i,(h,b) in enumerate(tests):
    col=i%3; row=i//3; card(s,.75+col*4.05,1.42+row*1.72,3.72,1.40,h,b,[ACCENT,BLUE,AMBER][col],f"{i+1:02d}")
add_rect(s,.75,5.35,11.85,.70,MINT,True,ACCENT)
add_text(s,"Plan → Prepare data → Execute → Capture evidence → Log defect → Fix → Retest → Regression",.95,5.60,11.45,.20,11.3,GREEN,True,PP_ALIGN.CENTER)
pill(s,"TEST EXECUTION: NOT STARTED",4.77,6.36,3.80,RED)

# 15 cases
s = base_slide("Test-Case Design", "Initial cases — keep status ‘Not Run’ until actual execution")
cases=[["ID","Scenario","Expected Result","Status"],
       ["TC-01","Invalid login","Access denied; generic error shown","Not Run"],
       ["TC-02","Select occupied bed","Allocation rejected","Not Run"],
       ["TC-03","Valid test payment","Payment and receipt recorded","Not Run"],
       ["TC-04","Student opens admin route","Access denied","Not Run"],
       ["TC-05","Expired/unapproved QR","Gate event rejected","Not Run"],
       ["TC-06","Resolve complaint","Status history retained","Not Run"],
       ["TC-07","Mobile dashboard","No overflow; actions usable","Not Run"]]
add_table(s,cases,.75,1.42,11.85,4.68,[.85,3.70,5.35,1.95],9.5)
add_text(s,"Record: test data · actual result · evidence · tester · date · pass/fail · defect ID",.78,6.36,11.55,.25,10,BODY,False,PP_ALIGN.CENTER)

# 16 defects/system testing
s = base_slide("Defect Report & System Testing", "Transparent tracking from issue discovery to verified closure")
add_text(s,"DEFECT RECORD",.75,1.42,2.0,.22,11.5,GREEN,True)
defect=[["Field","Required detail"],["Identity","Defect ID, module, build and environment"],
        ["Evidence","Steps, expected vs actual result, screenshot"],["Priority","Severity and business impact"],
        ["Closure","Fix version, retest result and tester"]]
add_table(s,defect,.75,1.82,5.65,3.30,[1.60,4.05],9.7)
add_text(s,"END-TO-END SYSTEM SCENARIOS",6.78,1.42,3.1,.22,11.5,GREEN,True)
add_bullets(s,["Admission → verification → allocation","Invoice → test payment → receipt","Complaint → assignment → resolution","Leave approval → QR → exit/re-entry logs","Unauthorized action → access denied","Backup → restore → record verification"],6.78,1.82,5.45,3.50,12)
pill(s,"NO DEFECTS RECORDED — TESTING NOT STARTED",3.95,6.10,5.42,AMBER)

# 17 demonstration
s = base_slide("Demonstration", "Planned live walkthrough after minimum viable product completion")
screenshot(s,ASSETS["demo"],.75,1.42,6.70,4.85,"Insert a real working-system screenshot")
steps=["Student login and profile","Submit hostel application","Warden document approval","Available-bed allocation","Invoice and test receipt","Leave approval and QR scan","Complaint assignment and closure"]
for i,text in enumerate(steps):
    y=1.43+i*.66
    add_rect(s,7.85,y,.42,.42,ACCENT,True)
    add_text(s,f"{i+1}",7.85,y+.13,.42,.14,8.5,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,text,8.48,y+.10,3.90,.24,10.8,INK,True)
pill(s,"DEMO PENDING IMPLEMENTATION",7.85,6.26,4.50,RED)

# 18 verification
s = base_slide("Verification & Performance", "Acceptance criteria — replace targets with measured evidence")
metrics=[["Area","Acceptance Criterion","Evidence"],
         ["Correctness","Critical workflows complete correctly","Executed system tests"],
         ["Authorization","Protected routes reject wrong roles","Access-control tests"],
         ["Integrity","No conflicting bed allocation","Constraint/transaction tests"],
         ["QR control","Only valid approved QR creates event","Valid/invalid QR tests"],
         ["Compatibility","Core pages work on target browsers","Browser test matrix"],
         ["Performance","Acceptable local response under basic load","Median and P95 timings"],
         ["Reliability","Backup restores complete data","Restore verification"]]
add_table(s,metrics,.75,1.42,11.85,4.78,[2.05,5.40,4.40],9.2)
add_text(s,"Do not invent timings or pass percentages. Measure after implementation and show test evidence.",.78,6.40,11.50,.25,10,RED,True,PP_ALIGN.CENTER)

# 19 plan
s = base_slide("Development Roadmap", "Balanced twelve-week route from design to demonstration")
phases=[("W1–2","Design Freeze","Stack, schema, UML"),("W3–4","Foundation","Auth, roles, base UI"),
        ("W5–6","Core","Admission, beds, allocation"),("W7–8","Services","Fees, complaints, leave"),
        ("W9–10","Integration","QR, reports, security"),("W11–12","Verification","Testing, fixes, demo")]
for i,(week,name,work) in enumerate(phases):
    x=.60+i*2.10; col=[ACCENT,BLUE,AMBER][i%3]
    add_rect(s,x,2.05,1.78,2.38,PALE,True,col)
    add_rect(s,x+.48,1.65,.82,.42,col,True)
    add_text(s,week,x+.48,1.78,.82,.14,8.5,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,name,x+.12,2.48,1.54,.24,11,col,True,PP_ALIGN.CENTER)
    add_text(s,work,x+.18,3.10,1.42,.55,9.5,BODY,False,PP_ALIGN.CENTER)
    if i<5: arrow(s,x+1.78,3.22,x+2.05,3.22,LINE,1.5)
add_text(s,"Execution principles",.75,5.10,2.0,.22,11.5,GREEN,True)
add_bullets(s,["Build the minimum viable product first","Assign one owner to every module","Integrate and review every week","Capture screenshots, tests and defects continuously"],.78,5.50,11.2,1.05,11.5)

# 20 conclusion
s = base_slide("Conclusion", "A credible Seminar–III outcome requires real implementation evidence")
add_rect(s,.75,1.42,11.85,1.10,GREEN,True)
add_text(s,"HOSTEL MONARCH is designed to connect hostel admission, allocation, finance, service and gate-security workflows in one role-based platform.",1.05,1.76,11.25,.40,15,WHITE,True,PP_ALIGN.CENTER)
summary=[("DESIGN","Requirements and workflows prepared",ACCENT),
         ("BUILD","Develop the focused core system",BLUE),
         ("TEST","Record genuine test and defect evidence",AMBER),
         ("DEMONSTRATE","Show actual workflows and limitations",GREEN)]
for i,(h,b,col) in enumerate(summary): card(s,.75+i*3.02,3.05,2.72,1.60,h,b,col,f"{i+1:02d}")
add_text(s,"Next milestone: confirm the stack and complete authentication, database schema and base dashboards.",.75,5.65,11.85,.38,13,INK,True,PP_ALIGN.CENTER)

# 21 Thank you
s = prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb = GREEN
add_rect(s,0,0,.22,7.5,ACCENT)
add_text(s,"THANK YOU",.95,1.40,8.0,.70,34,WHITE,True)
add_text(s,"Questions and Technical Suggestions are Welcome",.95,2.30,9.0,.32,16,WHITE)
add_rect(s,.95,2.90,2.20,.04,ACCENT)
add_text(s,"PROJECT TEAM",.95,3.42,2.0,.22,10,ACCENT,True)
add_text(s,"\n".join(f"{n} — {e}" for n,e in TEAM),.95,3.82,5.1,1.30,12.5,WHITE)
add_text(s,f"{PROJECT['college']}\n{PROJECT['department']}\n{PROJECT['course']}",7.10,3.42,4.9,1.20,13,WHITE)
add_text(s,"Guided by: "+PROJECT["guide"],7.10,5.10,4.0,.24,10.5,ACCENT,True)

prs.core_properties.title = "HOSTEL MONARCH — Seminar III Redesigned"
prs.core_properties.subject = "GTU Diploma Minor Project"
prs.core_properties.author = "Hostel Monarch Project Team"
prs.core_properties.comments = "Rudra Bhanderi removed. Planned work is explicitly labelled."
prs.save(OUTPUT)
print(f"Created {OUTPUT} with {len(prs.slides)} slides")